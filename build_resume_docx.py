from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT_FULLSTACK = "林小龙-全栈开发工程师简历-母版.docx"
OUT_FRONTEND = "林小龙-前端开发工程师简历-母版.docx"

FONT_CN = "微软雅黑"
FONT_EN = "Arial"
COLOR_INK = RGBColor(31, 41, 55)
COLOR_MUTED = RGBColor(91, 101, 117)
COLOR_ACCENT = RGBColor(29, 78, 216)
COLOR_LINE = "D7DEE8"


PROFILES = {
    "fullstack": {
        "out": OUT_FULLSTACK,
        "subtitle": "资深前端偏全栈开发工程师 / 全栈开发工程师",
        "advantages": [
            "7年以上 Web 应用开发经验，以复杂前端业务和 Electron 客户端交付为主，具备前端偏全栈开发、接口联调、容器化部署协作和问题排查能力。",
            "长期负责 PC 客户端、后台管理、H5 嵌入页、可视化平台等业务，能参与需求评审、技术方案拆解、核心模块设计、开发落地和迭代交付。",
            "具备复杂桌面端交互、IM 通讯、TRTC/WebRTC 音视频、多窗口通信、文件同步、本地缓存、地图/大屏可视化等业务经验。",
            "具备前端工程化和项目规范建设经验，能抽象通用组件/指令/进程通信逻辑，支撑多人协作和长期项目维护。",
            "熟悉 Vue、React、Electron、TypeScript、Node.js、Express、Koa2、Docker、Nginx，了解 Kubernetes 基础资源和服务排查流程。",
        ],
        "skills": [
            "前端技术：熟悉 JavaScript / TypeScript、Vue、React、Element UI、vue-element-admin，具备组件化开发、状态管理、复杂表单/权限后台和多端适配经验。",
            "Electron 客户端：熟悉主进程/渲染进程通信、多窗口管理、窗口透明/点穿、截图、本地文件目录处理、注册表读取和安装状态识别等桌面端能力。",
            "服务端协作：熟悉 Node.js、Express、Koa2，可参与登录鉴权、基础数据查询、文件/资源代理、业务配置等接口开发与前后端联调。",
            "部署与排查：了解 Docker 镜像构建、Docker Compose、Nginx、Kubernetes 基础资源配置，可配合完成环境配置、服务发布、日志查看和问题定位。",
            "工程化能力：熟悉 Vite、Webpack、Gulp、ESLint、Prettier、Mock、Git/SVN，具备项目规范落地、公共逻辑抽象、构建配置和协作流程优化经验。",
            "复杂业务集成：具备 TRTC/WebRTC、IM 通讯、地图 API、ECharts 可视化、文件同步、本地缓存等复杂业务接入和联调经验。",
        ],
        "include_node_project": True,
        "current_extra_bullets": [
            "参与 Node 服务接口联调与部分业务接口开发，配合后端打通用户、权限、文件、会议状态等模块的数据链路。",
            "参与 Docker 镜像构建、环境变量配置、测试/生产环境部署与发布问题排查，协助处理容器日志和服务异常。",
        ],
    },
    "frontend": {
        "out": OUT_FRONTEND,
        "subtitle": "资深 Web 前端开发工程师 / 高级前端开发工程师",
        "advantages": [
            "7年以上 Web 前端开发经验，长期负责 Electron 客户端、后台管理、H5 嵌入页、小程序、可视化平台等多端业务交付。",
            "熟悉 Vue、React、TypeScript、Electron，能独立推进需求拆解、技术方案设计、模块开发、组件封装、联调和迭代上线。",
            "具备多窗口通信、IM 通讯、TRTC/WebRTC 音视频、文件同步、本地缓存、地图/大屏可视化等复杂业务实践经验。",
            "具备前端工程化、项目规范建设、复杂组件拆分和公共逻辑抽象能力，能支撑多人协作下的稳定迭代和长期维护。",
            "了解 Node.js 服务开发、Docker/Nginx 部署流程和基础日志排查，能配合完成前后端联调、环境配置和发布问题定位。",
        ],
        "skills": [
            "前端技术：熟悉 JavaScript / TypeScript、Vue、React、Element UI、vue-element-admin，具备组件化开发、复杂表单、权限后台和多端适配经验。",
            "Electron 客户端：熟悉主进程/渲染进程通信、多窗口管理、窗口透明/点穿、截图、本地文件目录处理、注册表读取和安装状态识别等桌面端能力。",
            "工程化与规范：熟悉 Vite、Webpack、Gulp、ESLint、Prettier、Mock、Git/SVN，具备项目规范落地、公共逻辑抽象、构建配置和协作流程优化经验。",
            "复杂业务集成：具备 TRTC/WebRTC、IM 通讯、地图 API、ECharts 可视化、文件同步、本地缓存等复杂业务接入和联调经验。",
            "多端与业务场景：具备 PC 客户端、后台管理、H5 嵌入页、小程序、uni-app、可视化大屏等多端业务开发和适配经验。",
            "后端与交付协作：了解 Node.js、Express、Koa2、Docker、Nginx、环境变量配置和基础日志排查，可配合完成接口联调和发布问题定位。",
        ],
        "include_node_project": False,
        "current_extra_bullets": [],
    },
}


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT_EN
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_border_bottom(paragraph, color=COLOR_LINE, size="8", space="6"):
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, 12, True, COLOR_ACCENT)
    set_paragraph_border_bottom(p, size="6", space="4")
    return p


def add_body(doc, text="", size=9.5, bold=False, color=COLOR_INK, after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_run_font(run, size, bold, color)
    return p


def add_kv_line(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    label_run = p.add_run(label)
    set_run_font(label_run, 9.5, True, COLOR_INK)
    value_run = p.add_run(value)
    set_run_font(value_run, 9.5, False, COLOR_INK)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.45 + level * 0.3)
    p.paragraph_format.first_line_indent = Cm(-0.22)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    set_run_font(run, 9.2, False, COLOR_INK)
    return p


def add_role(doc, date, title, org):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    r1 = p.add_run(date)
    set_run_font(r1, 9.5, True, COLOR_INK)
    r2 = p.add_run("  |  ")
    set_run_font(r2, 9.5, False, COLOR_MUTED)
    r3 = p.add_run(title)
    set_run_font(r3, 9.5, True, COLOR_INK)
    r4 = p.add_run("  |  ")
    set_run_font(r4, 9.5, False, COLOR_MUTED)
    r5 = p.add_run(org)
    set_run_font(r5, 9.5, False, COLOR_INK)


def add_project(doc, date, name, stack, intro, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    r1 = p.add_run(f"{date}  {name}")
    set_run_font(r1, 10, True, COLOR_INK)
    p2 = add_body(doc, intro, size=9.2, after=1.5)
    p2.paragraph_format.line_spacing = 1.05
    add_kv_line(doc, "技术栈：", stack)
    for item in bullets:
        add_bullet(doc, item)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_EN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = COLOR_INK

    for style_name in ["List Bullet", "List Paragraph"]:
        style = styles[style_name]
        style.font.name = FONT_EN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(9.2)


def build(profile):
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("林小龙")
    set_run_font(run, 20, True, COLOR_INK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(5)
    r = subtitle.add_run(profile["subtitle"])
    set_run_font(r, 10.5, False, COLOR_MUTED)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(7)
    r = contact.add_run("1996年生  |  现居厦门  |  13290731298  |  1847426505@qq.com")
    set_run_font(r, 9.3, False, COLOR_MUTED)

    add_heading(doc, "个人优势")
    for item in profile["advantages"]:
        add_bullet(doc, item)

    add_heading(doc, "专业技能")
    for item in profile["skills"]:
        add_bullet(doc, item)

    add_heading(doc, "工作经历")
    add_role(doc, "2022.06 - 至今", "Web 前端高级工程师", "厦门玄空科技有限公司")
    add_bullet(doc, "担任核心开发人员，负责 Electron 客户端核心能力、后台管理模块、音视频会议、IM 通讯和文件同步等复杂业务迭代。")
    add_bullet(doc, "参与需求评审、技术方案设计、复杂功能拆解、项目规范建设和跨端联调，推动核心模块稳定交付。")

    add_role(doc, "2021.04 - 2022.06", "Web 前端工程师", "厦门互啪智能科技有限公司")
    add_bullet(doc, "负责自研 PC 客户端办公应用、后台管理系统、H5 嵌入原生移动端等业务模块开发和版本迭代。")
    add_bullet(doc, "参与 Electron 桌面端交互、富文本扩展、项目规范建设、公共逻辑封装和项目结构优化。")

    add_role(doc, "2018.08 - 2021.01", "Web 前端工程师", "福建觅觉文化信息有限公司")
    add_bullet(doc, "负责公司官网、小程序、后台管理、可视化数据平台等前端项目开发与维护。")
    add_bullet(doc, "参与巨幕播放系统和可视化平台的功能重构、复杂交互扩展、数据可视化和多项目场景交付。")

    add_role(doc, "2018.01 - 2018.06", "Web 前端实习生", "福建深空信息有限公司")
    add_bullet(doc, "参与前端页面构建、交互实现、接口联调和问题修复，完成从实习到正式前端开发的基础能力积累。")

    add_heading(doc, "项目经历")
    current_project_bullets = [
        "作为核心开发参与客户端方案拆解和模块落地，负责多窗口通信、事件分发、窗口透明、点穿、画中画等 Electron 桌面端能力。",
        "抽象基于 type 的进程消息监听与发布机制，统一多窗口事件处理入口，降低跨窗口状态同步和业务判断的维护成本。",
        "对接腾讯 TRTC 音视频能力，完成单聊/多聊/多人会议、会议窗口、画面分享、订阅/取消订阅和会议状态同步等复杂交互。",
        *profile["current_extra_bullets"],
        "设计本地共享目录与文件状态校验逻辑，支持目录动态匹配、命名生成、文件复制同步和异常状态处理。",
        "参与后台管理功能迭代，覆盖 CG 审核、任务流程、员工考勤、人事招聘、财务薪资、人员权限和系统配置等模块。",
    ]
    add_project(
        doc,
        "2022.06 - 至今",
        "天迈管理系统",
        "Vue / Electron / Element UI / Axios / Mock / ESLint / Prettier / TRTC / jeecgboot",
        "项目包含 Electron 客户端与后台管理两部分，覆盖多窗口桌面端交互、IM 通讯、腾讯 TRTC 音视频会议、文件同步、CG 审核、任务流程及人员/招聘/薪资/权限等后台模块；本人主要负责客户端核心能力、复杂交互和后台业务模块迭代。",
        current_project_bullets,
    )

    add_project(
        doc,
        "外接项目",
        "澳门美容管理平台",
        "Vue2 / uni-app / JeecgBoot / Ant Design Vue / Vuex / Vue Router / Axios / ColorUI / 微信小程序",
        "面向澳门美容门店经营场景的多端业务管理平台，覆盖平台管理端、门店后台端、会员小程序端和员工小程序端，支持多角色登录、后台运营、会员消费、员工履约和 OA 协同办公等业务闭环；本人全程负责前端业务模块开发，并协调项目搭建、部署、迭代和多端联调。",
        [
            "担任项目前端负责人，负责前端工程搭建、环境配置、部署上线、版本迭代和跨端联调协调，推进项目从开发到交付落地。",
            "独立完成门店后台端、会员小程序端、员工小程序端核心前端业务开发，覆盖会员、预约、销售、人事考勤、OA 协同、营销、仓存、报表、消息推送等模块。",
            "梳理后台管理端、小程序会员端和员工端之间的数据关系，打通会员消费、员工服务履约、后台运营管理的多角色多端业务联动。",
            "负责 OA 相关前端模块建设，包括员工工作台、申请审批、考勤/请假/调更、早会、任务、日记和消息通知等内部协同流程。",
            "适配澳门本地化业务场景，支持简体/繁体/英语国际化、地图定位、短信、微信服务号/小程序消息等能力接入与联调。",
        ],
    )

    if profile["include_node_project"]:
        add_project(
            doc,
            "相关实践",
            "Node 服务与容器化部署协作",
            "Node.js / Express / Koa2 / Docker / Docker Compose / Kubernetes / Nginx / Git",
            "围绕业务系统前后端交付、接口服务、部署发布和环境维护开展的工程化实践，重点支撑开发、测试、生产环境的一致性和可排查性。",
            [
                "参与登录鉴权、基础数据查询、文件/资源代理、业务配置等 Node.js 接口开发，并完成前后端联调。",
                "维护不同环境下的接口地址、环境变量、构建参数和部署配置，减少环境差异导致的联调和发布问题。",
                "编写或维护 Dockerfile / compose 配置，完成前端静态资源、Node 服务及相关依赖服务的容器化构建与启动。",
                "参与 Kubernetes 环境下的部署协作，理解 Deployment、Service、ConfigMap、Secret、Ingress 等基础资源配置。",
                "配合排查测试/生产环境问题，通过容器日志、服务状态、接口响应和网络配置定位发布失败、接口异常等问题。",
            ],
        )

    add_project(
        doc,
        "2021.04 - 2022.06",
        "MEUP 效率办公平台",
        "Vue / Electron / Element UI / Axios / Mock / ESLint / Prettier",
        "面向内部效率办公场景的平台，包含应用商城、工作台、知识库、悬赏广场、可视化平台、邮箱等模块。",
        [
            "负责多个客户端交互模块开发，覆盖全局可拖动便签、自定义右键弹窗、拖拽调整窗口区域、一键截图等能力。",
            "实现自定义 loading 指令，按页面接口请求状态统一展示加载反馈，改善弱网和多接口并发场景下的体验一致性。",
            "对 WangEditor 进行二次扩展，支持自定义菜单和业务插件接入，提升知识库/富文本场景的可扩展性。",
            "参与 Electron 桌面端能力建设，包括安装包界面优化、大分辨率 Icon 提取、应用安装状态识别和注册表读取。",
        ],
    )

    add_project(
        doc,
        "2020.03 - 2021.01",
        "新时代党建融平台",
        "Vue CLI / Electron / Cordova / ECharts / 百度地图 API / WebRTC / vue-element-admin",
        "面向党建机关单位的信息化平台，包含管理平台与可视化数据平台，支持党建资讯、党建视频、全景党建馆、党建会议、党建地图和主题风格等模块。",
        [
            "参与项目前两期核心开发，完成党建组织、资讯、视频会议、党建地图等多方数据与业务模块接入。",
            "基于 ECharts 和百度地图 API 实现组织数据可视化、平台分布展示、地图交互和功能操作入口。",
            "对接第三方音视频 SDK，基于 WebRTC 协议落地多应用办公场景下的视频会议能力。",
            "参与 vue-element-admin 后台管理平台建设，负责账号层级、权限分配、信息发布和可视化平台布局管理等模块。",
        ],
    )

    add_project(
        doc,
        "2018.10 - 2020.05",
        "巨幕播放系统",
        "Vue / ES6 / Sass / JavaScript / WPF WebView / UDP / ESLint / Git",
        "面向本地播放器和大屏展示场景的播控系统，基于 HTML 嵌入 WPF 本地播放器，支持视频、图片、序列帧、网页特效、弹幕、节目编排和移动端控制。",
        [
            "负责大屏播放场景核心交互开发，包括布局编排、播放区域拖拽定位、素材拉伸缩放和多媒体嵌套展示。",
            "通过网页场景嵌入播放器能力，实现音频可视化、人机互动、网页特效、账号管理等扩展功能。",
            "参与手机端/平板端中控模块开发，通过 UDP 或接口指令实现场景切换、图片/视频/弹幕上传和播放状态控制。",
            "在多个项目环境中持续扩展和维护系统能力，提升老项目的兼容性、稳定性和复用能力。",
        ],
    )

    add_heading(doc, "教育经历")
    add_body(doc, "2015.09 - 2018.06  |  福建农业职业技术学院  |  软件开发  |  大专", size=9.5, after=2)

    doc.save(profile["out"])
    return profile["out"]


if __name__ == "__main__":
    for profile in PROFILES.values():
        print(build(profile))
