from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT_FRONTEND = "林小龙-前端开发工程师简历-优化版.docx"
OUT_FULLSTACK = "林小龙-前端偏全栈开发工程师简历-优化版.docx"

FONT_CN = "微软雅黑"
FONT_EN = "Arial"
COLOR_INK = RGBColor(31, 41, 55)
COLOR_MUTED = RGBColor(90, 99, 115)
COLOR_ACCENT = RGBColor(20, 83, 150)
COLOR_LINE = "D7DEE8"


BASE = {
    "name": "林小龙",
    "contact": "30岁  |  现居厦门  |  13290731298  |  1847426505@qq.com",
    "education": "2015.09 - 2018.06  |  福建农业职业技术学院  |  软件开发  |  大专",
}


PROFILES = {
    "frontend": {
        "out": OUT_FRONTEND,
        "title": "高级 Web 前端开发工程师 / Electron 开发工程师",
        "summary": [
            "7年以上前端开发经验，主攻 Electron 客户端、复杂后台系统和多端业务交付，长期负责多窗口通信、音视频会议、文件同步、权限后台等复杂模块从方案拆解到上线落地。",
            "长期处理多窗口通信、TRTC/WebRTC 音视频、IM 通信、文件同步、本地缓存、地图与大屏可视化等复杂交互，擅长把业务流程拆成稳定可维护的模块。",
            "熟悉 Vue、React、TypeScript、Electron、Element UI、Ant Design Vue、uni-app，具备组件封装、状态管理、复杂表单、权限系统、多端适配和老项目持续迭代经验。",
            "具备工程化与协作规范意识，能够沉淀公共组件、指令、进程通信机制和项目约定，支撑多人协作、跨端联调与长期维护。",
        ],
        "skills": [
            "前端开发：JavaScript / TypeScript、Vue、React、Element UI、Ant Design Vue、vue-element-admin，熟悉组件化、状态管理、复杂表单、权限后台和多端适配。",
            "Electron 桌面端：熟悉主进程/渲染进程通信、多窗口状态同步、窗口透明/点穿、截图、注册表读取、安装状态识别、本地文件与目录处理等能力。",
            "复杂业务集成：具备 TRTC/WebRTC、IM 通信、地图 API、ECharts 可视化、文件同步、本地缓存、微信小程序消息等接入与联调经验。",
            "多端应用交付：熟悉 uni-app 小程序与 App 端开发，具备 HBuilderX 云打包/本地打包、App 权限配置、真机调试、接口联调和多端兼容处理经验。",
            "工程化质量：熟悉 Vite、Webpack、Gulp、ESLint、Prettier、Mock、Git/SVN，能够推进代码规范、公共逻辑抽象、构建配置和协作流程优化。",
            "交付协作：了解 Node.js、Express、Koa2、Docker、Nginx 与基础日志排查，可配合完成接口联调、环境配置、部署发布和问题定位。",
        ],
        "include_node_practice": False,
        "current_extra": [],
    },
    "fullstack": {
        "out": OUT_FULLSTACK,
        "title": "高级前端偏全栈开发工程师 / Electron 开发工程师",
        "summary": [
            "7年以上 Web 应用开发经验，以复杂前端、Electron 客户端和业务系统交付为主，具备 Node 接口协作、容器化部署配合和线上问题排查经验。",
            "熟悉 Vue、TypeScript、Electron、Node.js、Express、Koa2、Docker、Nginx，能够参与前端核心模块、基础接口、环境配置和发布链路落地。",
            "长期负责 PC 客户端、后台管理、小程序、H5 嵌入页和可视化平台，能从需求评审、方案拆解、模块开发、跨端联调到迭代发布持续推进。",
            "擅长处理多窗口通信、IM 通信、TRTC/WebRTC 音视频、文件同步、本地缓存、地图与大屏可视化等复杂业务，并沉淀可复用的组件、指令和通信机制。",
        ],
        "skills": [
            "前端技术：JavaScript / TypeScript、Vue、React、Element UI、Ant Design Vue、vue-element-admin，熟悉组件化、复杂表单、权限后台和多端适配。",
            "Electron 桌面端：熟悉主进程/渲染进程通信、多窗口管理、窗口透明/点穿、截图、注册表读取、安装状态识别、本地文件与目录处理等能力。",
            "Node 与接口协作：熟悉 Node.js、Express、Koa2，可参与登录鉴权、基础数据查询、文件/资源代理、业务配置等接口开发与前后端联调。",
            "部署与排查：了解 Docker、Docker Compose、Nginx、Kubernetes 基础资源配置，能配合完成环境变量配置、服务发布、日志查看和问题定位。",
            "工程化能力：熟悉 Vite、Webpack、Gulp、ESLint、Prettier、Mock、Git/SVN，具备项目规范落地、公共逻辑抽象和协作流程优化经验。",
        ],
        "include_node_practice": True,
        "current_extra": [
            "参与 Node 服务接口联调与部分业务接口开发，配合打通用户、权限、文件、会议状态等模块的数据链路。",
            "参与 Docker 镜像构建、环境变量配置、测试/生产环境部署与发布问题排查，协助定位容器日志和服务异常。",
        ],
    },
}


WORK_EXPERIENCE = [
    (
        "2022.06 - 至今",
        "Web 前端高级工程师",
        "厦门玄空科技有限公司",
        [
            "担任核心开发人员，负责 Electron 客户端核心能力、后台管理模块、音视频会议、IM 通信和文件同步等复杂业务。",
            "参与需求评审、技术方案设计、功能拆解、项目规范建设和跨端联调，推动客户端与后台模块稳定迭代。",
        ],
    ),
    (
        "2021.04 - 2022.06",
        "Web 前端工程师",
        "厦门互啪智能科技有限公司",
        [
            "负责自研 PC 客户端办公应用、后台管理系统、H5 嵌入原生移动端等模块开发和版本迭代。",
            "参与 Electron 桌面端交互、富文本扩展、公共逻辑封装、项目规范建设和项目结构优化。",
        ],
    ),
    (
        "2018.08 - 2021.01",
        "Web 前端工程师",
        "福建觅觉文化信息有限公司",
        [
            "负责公司官网、小程序、后台管理、可视化数据平台等前端项目开发与维护。",
            "参与巨幕播放系统和可视化平台重构，持续扩展复杂交互、数据可视化和多项目交付能力。",
        ],
    ),
    (
        "2018.01 - 2018.06",
        "Web 前端实习生",
        "福建深空信息有限公司",
        ["参与前端页面构建、交互实现、接口联调和问题修复，完成从实习到正式前端开发的基础能力积累。"],
    ),
]


def current_project_bullets(extra):
    return [
        "围绕 Electron 客户端多窗口协同场景，建设多窗口通信、事件分发、窗口透明、点穿、画中画等能力，支撑会议、文件、通知等模块的跨窗口状态同步。",
        "抽象基于 type 的进程消息监听与发布机制，统一多窗口事件入口，降低跨窗口状态同步和业务判断维护成本。",
        "对接腾讯 TRTC 音视频能力，完成单聊、多聊、多人会议、画面分享、订阅/取消订阅、会议状态同步等交互，配合 IM 与客户端窗口完成会议链路闭环。",
        *extra,
        "设计本地共享目录与文件状态校验逻辑，支持目录动态匹配、命名生成、文件复制同步和异常状态处理。",
        "参与后台管理模块迭代，覆盖 CG 审核、任务流程、员工考勤、人事招聘、财务薪资、人员权限和系统配置。",
    ]


PROJECTS = [
    {
        "date": "2022.06 - 至今",
        "name": "天迈管理系统",
        "stack": "Vue / Electron / Element UI / Axios / Mock / ESLint / Prettier / TRTC / JeecgBoot",
        "intro": "项目包含 Electron 客户端与后台管理两部分，覆盖桌面端多窗口交互、IM 通信、TRTC 音视频会议、文件同步、CG 审核、任务流程及人员/招聘/薪资/权限等模块。",
        "bullets": current_project_bullets,
    },
    {
        "date": "外接项目",
        "name": "澳门美容管理平台",
        "stack": "Vue2 / uni-app / App 打包 / JeecgBoot / Ant Design Vue / Vuex / Vue Router / Axios / ColorUI / 微信小程序",
        "intro": "面向澳门美容门店经营场景的多端业务管理平台，覆盖平台管理端、门店后台端、会员小程序端和员工小程序端，支持多角色登录、会员消费、员工履约、后台运营和 OA 协同。",
        "bullets": [
            "担任项目前端负责人，负责工程搭建、环境配置、部署上线、版本迭代和跨端联调，独立完成门店后台、会员端、员工端核心业务开发。",
            "基于 uni-app 完成小程序与 App 端适配，参与 HBuilderX 打包、App 权限配置、真机调试和多端兼容处理，保障移动端交付落地。",
            "梳理后台管理端、会员端和员工端数据关系，打通会员消费、员工服务履约、后台运营管理的多角色多端联动。",
            "负责 OA 前端模块建设，包括员工工作台、申请审批、考勤/请假/调更、早会、任务、日记和消息通知等流程。",
            "适配澳门本地化场景，支持简体/繁体/英语国际化、地图定位、短信、微信服务号/小程序消息等能力接入。",
        ],
    },
    {
        "date": "2021.04 - 2022.06",
        "name": "MEUP 效率办公平台",
        "stack": "Vue / Electron / Element UI / Axios / Mock / ESLint / Prettier",
        "intro": "面向内部效率办公场景的平台，包含应用商城、工作台、知识库、悬赏广场、可视化平台、邮箱等模块。",
        "bullets": [
            "负责客户端交互模块开发，覆盖全局可拖动便签、自定义右键弹窗、拖拽调整窗口区域、一键截图等能力。",
            "实现自定义 loading 指令，按接口请求状态统一展示加载反馈，改善弱网和多接口并发场景下的体验一致性。",
            "对 WangEditor 进行二次扩展，支持自定义菜单和业务插件接入，提升知识库/富文本场景扩展能力。",
            "参与 Electron 桌面端能力建设，包括安装包界面优化、大分辨率 Icon 提取、应用安装状态识别和注册表读取。",
        ],
    },
    {
        "date": "2020.03 - 2021.01",
        "name": "新时代党建融平台",
        "stack": "Vue CLI / Electron / Cordova / ECharts / 百度地图 API / WebRTC / vue-element-admin",
        "intro": "面向党建机关单位的信息化平台，包含管理平台与可视化数据平台，支持党建资讯、党建视频、全景党建馆、党建会议、党建地图和主题风格等模块。",
        "bullets": [
            "参与项目前两期核心开发，完成党建组织、资讯、视频会议、党建地图等多方数据与业务模块接入。",
            "基于 ECharts 和百度地图 API 实现组织数据可视化、平台分布展示、地图交互和功能操作入口。",
            "对接第三方音视频 SDK，基于 WebRTC 协议落地多应用办公场景下的视频会议能力。",
            "参与 vue-element-admin 后台建设，负责账号层级、权限分配、信息发布和可视化平台布局管理等模块。",
        ],
    },
    {
        "date": "2018.10 - 2020.05",
        "name": "巨幕播放系统",
        "stack": "Vue / ES6 / Sass / JavaScript / WPF WebView / UDP / ESLint / Git",
        "intro": "面向本地播放器和大屏展示场景的播控系统，基于 HTML 嵌入 WPF 本地播放器，支持视频、图片、序列帧、网页特效、弹幕、节目编排和移动端控制。",
        "bullets": [
            "负责大屏播放场景核心交互开发，包括布局编排、播放区域拖拽定位、素材拉伸缩放和多媒体嵌套展示。",
            "通过网页场景嵌入播放器能力，实现音频可视化、人机互动、网页特效、账号管理等扩展功能。",
            "参与手机端/平板端中控模块开发，通过 UDP 或接口指令实现场景切换、素材上传和播放状态控制。",
            "在多个项目环境中持续扩展和维护系统能力，提升老项目兼容性、稳定性和复用能力。",
        ],
    },
]


NODE_PRACTICE = {
    "date": "相关实践",
    "name": "Node 服务与容器化部署协作",
    "stack": "Node.js / Express / Koa2 / Docker / Docker Compose / Kubernetes / Nginx / Git",
    "intro": "围绕业务系统前后端交付、接口服务、部署发布和环境维护开展工程化实践，重点支撑开发、测试、生产环境的一致性和可排查性。",
    "bullets": [
        "参与登录鉴权、基础数据查询、文件/资源代理、业务配置等 Node.js 接口开发，并完成前后端联调。",
        "维护接口地址、环境变量、构建参数和部署配置，减少环境差异导致的联调和发布问题。",
        "编写或维护 Dockerfile / compose 配置，完成前端静态资源、Node 服务及依赖服务的容器化构建与启动。",
        "参与 Kubernetes 部署协作，理解 Deployment、Service、ConfigMap、Secret、Ingress 等基础资源配置。",
        "配合排查测试/生产环境问题，通过容器日志、服务状态、接口响应和网络配置定位发布失败、接口异常等问题。",
    ],
}


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT_EN
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_border_bottom(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), COLOR_LINE)
    p_bdr.append(bottom)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.15)
    section.bottom_margin = Cm(1.15)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_EN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = COLOR_INK

    for style_name in ["List Bullet", "List Paragraph"]:
        style = doc.styles[style_name]
        style.font.name = FONT_EN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(9.0)


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3.5)
    run = p.add_run(text)
    set_run_font(run, 11.5, True, COLOR_ACCENT)
    set_paragraph_border_bottom(p)
    return p


def add_body(doc, text, size=9.2, bold=False, color=COLOR_INK, after=1.8, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.03
    run = p.add_run(text)
    set_run_font(run, size, bold, color)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.42 + level * 0.28)
    p.paragraph_format.first_line_indent = Cm(-0.2)
    p.paragraph_format.space_after = Pt(1.05)
    p.paragraph_format.line_spacing = 1.02
    run = p.add_run(text)
    set_run_font(run, 9.0, False, COLOR_INK)
    return p


def add_role(doc, date, title, org, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    for text, bold, color in [
        (date, True, COLOR_INK),
        ("  |  ", False, COLOR_MUTED),
        (title, True, COLOR_INK),
        ("  |  ", False, COLOR_MUTED),
        (org, False, COLOR_INK),
    ]:
        run = p.add_run(text)
        set_run_font(run, 9.2, bold, color)
    for item in bullets:
        add_bullet(doc, item)


def add_project(doc, project, extra=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3.5)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(f"{project['date']}  {project['name']}")
    set_run_font(run, 9.6, True, COLOR_INK)

    add_body(doc, project["intro"], size=8.95, after=1.0)
    add_body(doc, f"技术栈：{project['stack']}", size=8.95, bold=True, color=COLOR_MUTED, after=1.0)

    bullets = project["bullets"](extra or []) if callable(project["bullets"]) else project["bullets"]
    for item in bullets:
        add_bullet(doc, item)


def build(profile):
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run(BASE["name"])
    set_run_font(run, 20, True, COLOR_INK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(4)
    run = subtitle.add_run(profile["title"])
    set_run_font(run, 10.1, False, COLOR_MUTED)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(6)
    run = contact.add_run(BASE["contact"])
    set_run_font(run, 9.0, False, COLOR_MUTED)

    add_heading(doc, "职业定位")
    for item in profile["summary"]:
        add_bullet(doc, item)

    add_heading(doc, "专业技能")
    for item in profile["skills"]:
        add_bullet(doc, item)

    add_heading(doc, "工作经历")
    for role in WORK_EXPERIENCE:
        add_role(doc, *role)

    add_heading(doc, "项目经历")
    add_project(doc, PROJECTS[0], profile["current_extra"])
    add_project(doc, PROJECTS[1])
    if profile["include_node_practice"]:
        add_project(doc, NODE_PRACTICE)
    for project in PROJECTS[2:]:
        add_project(doc, project)

    add_heading(doc, "教育经历")
    add_body(doc, BASE["education"], size=9.2, after=2)

    doc.save(profile["out"])
    return profile["out"]


if __name__ == "__main__":
    for profile in PROFILES.values():
        print(build(profile))
