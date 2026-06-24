from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = "林小龙-资深前端开发工程师简历-微调版.docx"

FONT_CN = "微软雅黑"
FONT_EN = "Arial"
COLOR_INK = RGBColor(31, 41, 55)
COLOR_MUTED = RGBColor(86, 96, 112)
COLOR_ACCENT = RGBColor(20, 83, 150)
COLOR_RULE = "D7DEE8"


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT_EN
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_bottom_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), COLOR_RULE)
    p_bdr.append(bottom)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.12)
    section.bottom_margin = Cm(1.05)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_EN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = COLOR_INK

    for style_name in ["List Bullet", "List Paragraph"]:
        style = doc.styles[style_name]
        style.font.name = FONT_EN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(9.0)


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_run_font(run, 11.5, True, COLOR_ACCENT)
    add_bottom_border(p)
    return p


def add_body(doc, text, size=9.05, bold=False, color=COLOR_INK, before=0, after=1.8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.04
    run = p.add_run(text)
    set_run_font(run, size, bold, color)
    return p


def add_bullet(doc, text, level=0, after=1.0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.43 + level * 0.28)
    p.paragraph_format.first_line_indent = Cm(-0.2)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.02
    run = p.add_run(text)
    set_run_font(run, 8.95, False, COLOR_INK)
    return p


def add_inline_label(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.2)
    p.paragraph_format.line_spacing = 1.02
    r1 = p.add_run(label)
    set_run_font(r1, 9.0, True, COLOR_INK)
    r2 = p.add_run(value)
    set_run_font(r2, 9.0, False, COLOR_INK)


def add_role(doc, date, title, org, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    for text, bold, color in [
        (date, True, COLOR_INK),
        ("  |  ", False, COLOR_MUTED),
        (title, True, COLOR_INK),
        ("  |  ", False, COLOR_MUTED),
        (org, False, COLOR_INK),
    ]:
        run = p.add_run(text)
        set_run_font(run, 9.15, bold, color)
    for item in bullets:
        add_bullet(doc, item)


def add_project(doc, date, name, stack, intro, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3.4)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(f"{date}  {name}")
    set_run_font(run, 9.55, True, COLOR_INK)
    add_body(doc, intro, size=8.85, after=0.9)
    add_inline_label(doc, "技术栈：", stack)
    for item in bullets:
        add_bullet(doc, item, after=0.85)


def build():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(1.5)
    run = title.add_run("林小龙")
    set_run_font(run, 20, True, COLOR_INK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(3.5)
    run = subtitle.add_run("资深前端开发工程师 / 前端架构与 AI Native 工程化方向")
    set_run_font(run, 10.2, False, COLOR_MUTED)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(5.5)
    run = contact.add_run("1996年生  |  现居厦门  |  13290731298  |  1847426505@qq.com")
    set_run_font(run, 9.0, False, COLOR_MUTED)

    add_heading(doc, "职业定位")
    for item in [
        "7年以上前端开发经验，定位资深前端工程师，长期负责复杂中后台、Electron 桌面端、多端业务和可视化项目交付。",
        "熟悉 Vue / React / TypeScript / Electron / Node.js 基础能力，能从需求拆解、技术方案、组件边界、状态管理、接口契约到测试上线完整推进。",
        "具备复杂业务实现经验：多窗口通信、IM 通讯、TRTC/WebRTC 音视频、文件同步、本地缓存、权限后台、ECharts/地图可视化、H5/小程序/uni-app 多端适配。",
        "具备工程化与长期维护意识，能沉淀公共组件、指令、进程通信机制、开发规范和问题排查文档，支撑多人协作和存量系统持续演进。",
        "熟悉 AI 辅助开发工作流，能使用 Codex / Cursor / Claude Code / Copilot 等工具拆解任务、生成方案、审查产出，并结合人工判断补齐性能、交互、可维护性等 AI 盲区。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "专业技能")
    for item in [
        "前端框架：熟悉 Vue2/Vue3、React、JavaScript/TypeScript、Vuex/Pinia、React Hooks、Vue Router、Element UI / Element Plus、Ant Design Vue、vue-element-admin。",
        "工程化体系：熟悉 Vite、Webpack、Gulp、npm/pnpm、ESLint、Prettier、Mock、Git/SVN，可落地目录规范、代码规范、构建配置、组件复用和协作流程优化。",
        "架构与复杂模块：具备中后台权限、复杂表单、状态管理、接口封装、异常处理、数据看板、组件库/公共工具封装、多端兼容和存量系统重构经验。",
        "Electron 桌面端：熟悉主进程/渲染进程通信、多窗口管理、窗口透明/点穿、截图、注册表读取、安装状态识别、本地文件目录处理和跨窗口状态同步。",
        "性能与质量：能使用 DevTools 分析首屏加载、长任务、重排重绘、资源加载和内存问题，具备 Web 性能优化、兼容性处理、自动化测试策略和线上问题排查能力。",
        "AI Native 协作：能将需求拆成 AI 可执行任务，审查 AI 生成代码/文档/测试用例，使用 Playwright 等工具辅助回归验证，并把高频问题沉淀成规范或脚手架约束。",
        "服务端与交付协作：熟悉 Node.js、Express、Koa2 基础使用，了解 BFF/Mock 服务、接口聚合、Docker、Nginx、环境变量配置、日志查看和发布问题定位。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "工作经历")
    add_role(
        doc,
        "2022.06 - 至今",
        "Web 前端高级工程师",
        "厦门玄空科技有限公司",
        [
            "作为核心开发人员，负责 Electron 客户端核心能力、后台管理模块、音视频会议、IM 通讯、文件同步和权限业务等复杂模块迭代。",
            "参与需求评审、技术方案设计、模块拆解、接口联调、问题排查和项目规范建设，推动客户端与后台系统稳定交付。",
            "在多人协作和存量系统迭代场景中，持续抽象公共逻辑、统一事件通信和开发约定，降低后续需求接入成本。",
        ],
    )
    add_role(
        doc,
        "2021.04 - 2022.06",
        "Web 前端工程师",
        "厦门互啪智能科技有限公司",
        [
            "负责自研 PC 客户端办公应用、后台管理系统、H5 嵌入原生移动端等模块开发和版本迭代。",
            "参与 Electron 桌面端交互、富文本扩展、公共逻辑封装、项目规范建设和项目结构优化。",
        ],
    )
    add_role(
        doc,
        "2018.08 - 2021.01",
        "Web 前端工程师",
        "福建觅觉文化信息有限公司",
        [
            "负责官网、小程序、后台管理、可视化数据平台、巨幕播放系统等前端项目开发与维护。",
            "参与可视化平台和播放系统重构，积累复杂交互、数据可视化、多项目并行交付和老系统维护经验。",
        ],
    )
    add_role(
        doc,
        "2018.01 - 2018.06",
        "Web 前端实习生",
        "福建深空信息有限公司",
        ["参与前端页面构建、交互实现、接口联调和问题修复，完成前端开发基础能力积累。"],
    )

    add_heading(doc, "项目经历")
    add_project(
        doc,
        "2022.06 - 至今",
        "天迈管理系统",
        "Vue / Electron / Element UI / Axios / Mock / ESLint / Prettier / TRTC / JeecgBoot",
        "项目包含 Electron 客户端与后台管理两部分，覆盖多窗口桌面端交互、IM 通讯、腾讯 TRTC 音视频会议、文件同步、CG 审核、任务流程、人员/招聘/薪资/权限等模块；本人负责客户端核心能力、复杂交互和后台业务模块迭代。",
        [
            "主导多个 Electron 客户端复杂能力落地，建设多窗口通信、事件分发、窗口透明、点穿、画中画等能力，支撑会议、文件、通知等跨窗口状态同步。",
            "抽象基于 type 的进程消息监听与发布机制，统一多窗口事件入口，降低业务模块接入和跨窗口状态维护成本。",
            "对接腾讯 TRTC 音视频能力，完成单聊、多聊、多人会议、画面分享、订阅/取消订阅、会议状态同步等交互，配合 IM 与客户端窗口完成会议链路闭环。",
            "设计本地共享目录与文件状态校验逻辑，支持目录动态匹配、命名生成、文件复制同步和异常状态处理。",
            "参与后台管理模块迭代，覆盖 CG 审核、任务流程、员工考勤、人事招聘、财务薪资、人员权限和系统配置等复杂中后台场景。",
        ],
    )
    add_project(
        doc,
        "2025.03 - 至今",
        "STO 中文学习平台",
        "React / TypeScript / Vite / React Router / Zustand / Axios / i18n / Vue3 / JEECG Boot / Ant Design Vue",
        "面向海外中文学习与运营增长场景的三端协同项目，包含官网展示端、App 内嵌 H5 页和后台管理端；后台负责课程、内容、活动、积分商品等配置管理，前台负责动态展示、分享转化和 App WebView 场景承接。",
        [
            "负责官网及 App 内嵌 H5 核心页面开发，基于 React + TypeScript + Vite 落地课程展示、活动页、分享落地页、商品详情、协议与在线客服等业务场景。",
            "设计 H5 路由与 App WebView 接入逻辑，统一处理 URL 透传的 token、device_type、lang 等上下文参数，支撑登录态识别、设备信息上报和多语言展示。",
            "封装前台请求层，统一注入 token、语言、设备信息等请求头，并对接 /common/app、/market/app、/study 等接口域，降低页面侧重复联调成本。",
            "参与 Vue3 + JEECG Boot 后台管理端开发，完成课程管理、活动配置、积分商品、分享配置、官网内容配置等模块的表单、列表和状态流转能力。",
            "配合后端打通“后台配置 - 接口下发 - 官网/H5 动态展示”链路，保障活动营销、分享转化和 App 内容运营可以通过后台配置快速调整。",
        ],
    )
    add_project(
        doc,
        "外接项目",
        "澳门美容管理平台",
        "Vue2 / uni-app / JeecgBoot / Ant Design Vue / Vuex / Vue Router / Axios / ColorUI / 微信小程序",
        "面向澳门美容门店经营场景的多端业务管理平台，覆盖平台管理端、门店后台端、会员小程序端和员工小程序端，支持多角色登录、会员消费、员工履约、后台运营和 OA 协同。",
        [
            "担任项目前端负责人，负责工程搭建、环境配置、部署上线、版本迭代和跨端联调，独立完成门店后台、会员端、员工端核心业务开发。",
            "基于 uni-app 完成小程序与 App 端适配，参与 HBuilderX 打包、App 权限配置、真机调试和多端兼容处理，保障移动端交付落地。",
            "梳理后台管理端、会员端和员工端数据关系，打通会员消费、员工服务履约、后台运营管理的多角色多端联动。",
            "负责 OA 前端模块建设，包括员工工作台、申请审批、考勤/请假/调更、早会、任务、日记和消息通知等流程。",
            "适配澳门本地化场景，支持简体/繁体/英语国际化、地图定位、短信、微信服务号/小程序消息等能力接入。",
        ],
    )
    add_project(
        doc,
        "AI Native 实践",
        "AI 辅助前端研发与质量把关",
        "Codex / Cursor / Claude Code / GitHub Copilot / Playwright / DevTools / Markdown",
        "围绕资深前端日常研发，将 AI 工具作为任务拆解、代码生成、文档沉淀、测试辅助和问题排查工具使用，重点保留人工在架构决策、质量审查、性能判断和交互体验上的把关。",
        [
            "将产品需求拆解为 AI 可执行的页面、组件、接口、测试和文档任务，明确输入、验收标准和边界条件，减少反复返工。",
            "审查 AI 生成代码的状态管理、组件边界、异常路径、性能风险和可维护性问题，避免盲目合并看似可运行但长期成本较高的实现。",
            "结合 Playwright 自动化脚本做页面回归、流程验证和重复操作辅助，同时保留人工对体感性能、复杂交互和 UI 细节的最终判断。",
            "将 AI 频繁出错的问题沉淀为开发规范、代码模板或检查清单，让后续 AI 产出更稳定，而不是把重复修补变成个人负担。",
        ],
    )
    add_project(
        doc,
        "2021.04 - 2022.06",
        "MEUP 效率办公平台",
        "Vue / Electron / Element UI / Axios / Mock / ESLint / Prettier",
        "面向内部效率办公场景的平台，包含应用商城、工作台、知识库、悬赏广场、可视化平台、邮箱等模块。",
        [
            "负责客户端交互模块开发，覆盖全局可拖动便签、自定义右键弹窗、拖拽调整窗口区域、一键截图等能力。",
            "实现自定义 loading 指令，按接口请求状态统一展示加载反馈，改善弱网和多接口并发场景下的体验一致性。",
            "对 WangEditor 进行二次扩展，支持自定义菜单和业务插件接入，提升知识库/富文本场景扩展能力。",
            "参与 Electron 桌面端能力建设，包括安装包界面优化、大分辨率 Icon 提取、应用安装状态识别和注册表读取。",
        ],
    )
    add_project(
        doc,
        "2020.03 - 2021.01",
        "新时代党建融平台",
        "Vue CLI / Electron / Cordova / ECharts / 百度地图 API / WebRTC / vue-element-admin",
        "面向党建机关单位的信息化平台，包含管理平台与可视化数据平台，支持党建资讯、党建视频、全景党建馆、党建会议、党建地图和主题风格等模块。",
        [
            "参与项目前两期核心开发，完成党建组织、资讯、视频会议、党建地图等多方数据与业务模块接入。",
            "基于 ECharts 和百度地图 API 实现组织数据可视化、平台分布展示、地图交互和功能操作入口。",
            "对接第三方音视频 SDK，基于 WebRTC 协议落地多应用办公场景下的视频会议能力。",
            "参与 vue-element-admin 后台建设，负责账号层级、权限分配、信息发布和可视化平台布局管理等模块。",
        ],
    )
    add_project(
        doc,
        "2018.10 - 2020.05",
        "巨幕播放系统",
        "Vue / ES6 / Sass / JavaScript / WPF WebView / UDP / ESLint / Git",
        "面向本地播放器和大屏展示场景的播控系统，基于 HTML 嵌入 WPF 本地播放器，支持视频、图片、序列帧、网页特效、弹幕、节目编排和移动端控制。",
        [
            "负责大屏播放场景核心交互开发，包括布局编排、播放区域拖拽定位、素材拉伸缩放和多媒体嵌套展示。",
            "通过网页场景嵌入播放器能力，实现音频可视化、人机互动、网页特效、账号管理等扩展功能。",
            "参与手机端/平板端中控模块开发，通过 UDP 或接口指令实现场景切换、素材上传和播放状态控制。",
            "在多个项目环境中持续扩展和维护系统能力，提升老项目兼容性、稳定性和复用能力。",
        ],
    )

    add_heading(doc, "教育经历")
    add_body(doc, "2015.09 - 2018.06  |  福建农业职业技术学院  |  软件开发  |  大专", size=9.05, after=2)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
